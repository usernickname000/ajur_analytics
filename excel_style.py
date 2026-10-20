# ============================================================
# excel_style.py — оформление Excel-отчётов и графики
# Общее для analytics (основной отчёт) и comparison (сравнение).
# ============================================================

import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Палитра (фирменный оранжевый + нейтральные) ─────────────
ACCENT      = 'F38120'
ACCENT_2    = '4A90D9'   # вторая серия на графиках
POSITIVE    = '22C55E'

HEADER_FILL = PatternFill('solid', fgColor=ACCENT)
ROW_ALT     = PatternFill('solid', fgColor='FFF4EC')
ROW_NORMAL  = PatternFill('solid', fgColor='FFFFFF')
RED_FILL    = PatternFill('solid', fgColor='FDECEA')
TOTAL_FILL  = PatternFill('solid', fgColor='FDE8D0')
SEP_FILL    = PatternFill('solid', fgColor='F0F0F0')
SIGNAL_FILLS = {'🔴': PatternFill('solid', fgColor='FDECEA'),
                '🟡': PatternFill('solid', fgColor='FEF9E7')}
SIGNAL_GRN  = PatternFill('solid', fgColor='EAFAF1')
NEW_FILL    = PatternFill('solid', fgColor='E2EFDA')   # зелёный — новые
LOST_FILL   = PatternFill('solid', fgColor='FCE4D6')   # красный — ушедшие

HEADER_FONT = Font(name='Segoe UI', bold=True, color='FFFFFF', size=10)
BODY_FONT   = Font(name='Segoe UI', size=9)
TOTAL_FONT  = Font(name='Segoe UI', bold=True, size=9)
SEP_FONT    = Font(name='Segoe UI', bold=True, color='888888', size=8)
NEG_FONT    = Font(name='Segoe UI', size=9, color='C0392B')
POS_FONT    = Font(name='Segoe UI', size=9, color='375623')
TITLE_FONT  = Font(name='Segoe UI', size=16, bold=True, color=ACCENT)

BORDER        = Border(bottom=Side(style='thin', color='E8E8E8'))
HEADER_ALIGN  = Alignment(horizontal='center', vertical='center', wrap_text=True)
BODY_ALIGN    = Alignment(vertical='center')

TOTAL_KEYWORDS  = {'итого', 'total', 'план', 'выполнение', 'grand', '═══', '──'}
SEPARATOR_CHARS = {'═', '─', '—'}


def style_header_row(ws, row_idx=1):
    for cell in ws[row_idx]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER


def style_body_cell(cell, fill=None, font=BODY_FONT):
    cell.font = font
    cell.border = BORDER
    cell.alignment = BODY_ALIGN
    if fill is not None:
        cell.fill = fill


def autosize_columns(ws, df, max_width=60):
    """Ширина колонок по самому длинному значению (DataFrame уже на листе)."""
    for col_idx, col in enumerate(df.columns, 1):
        try:
            col_max = df[col].dropna().astype(str).map(len).max()
            col_max = int(col_max) if pd.notna(col_max) else 10
            width = max(col_max, len(str(col))) + 2
        except Exception:
            width = 14
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, max_width)


def autosize_sheet(ws, min_w=10, max_w=50):
    """Ширина колонок по содержимому листа (когда DataFrame нет)."""
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, min_w), max_w)


def style_workbook(wb, log=print):
    """
    Применяет единое оформление ко всем листам книги:
    - оранжевые заголовки с белым текстом
    - чередующиеся строки (зебра)
    - красный фон для отрицательных чисел
    - жирный шрифт для итоговых строк
    - заморозка первой строки
    """
    try:
        for ws in wb.worksheets:
            if ws.title in ('📊 Графики',):
                continue
            max_row, max_col = ws.max_row, ws.max_column
            if max_row < 2:
                continue

            style_header_row(ws, 1)
            ws.row_dimensions[1].height = 22
            ws.freeze_panes = 'A2'
            is_signals = ws.title.startswith('00_Сигналы')

            for row in range(2, max_row + 1):
                first_val = str(ws.cell(row, 1).value or '').strip()
                is_separator = any(c in first_val for c in SEPARATOR_CHARS)
                is_total = any(k in first_val.lower() for k in TOTAL_KEYWORDS)

                if is_separator:
                    row_fill, row_font = SEP_FILL, SEP_FONT
                elif is_total:
                    row_fill, row_font = TOTAL_FILL, TOTAL_FONT
                elif is_signals:
                    row_fill = next((f for icon, f in SIGNAL_FILLS.items() if icon in first_val), SIGNAL_GRN)
                    row_font = BODY_FONT
                else:
                    row_fill = ROW_ALT if row % 2 == 0 else ROW_NORMAL
                    row_font = BODY_FONT

                for col in range(1, max_col + 1):
                    cell = ws.cell(row, col)
                    style_body_cell(cell, row_fill, row_font)
                    # Отрицательные числа — красный фон
                    if (not is_separator and not is_total
                            and isinstance(cell.value, (int, float)) and cell.value < 0):
                        cell.fill = RED_FILL
                        cell.font = NEG_FONT
                ws.row_dimensions[row].height = 18

        log("Оформление листов применено ✅")
    except Exception as e:
        log(f"⚠ Оформление не применено: {e}")


def _find_rev_col(ws, fallback=2):
    """Ищет в заголовке листа колонку с выручкой в тыс. руб."""
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val and 'тыс' in str(val).lower() and 'выруч' in str(val).lower():
            return c
    return fallback


def make_chart(kind, title, width=22, height=14, x_title=None, y_title=None):
    """kind: 'line' | 'bar' (горизонтальный) | 'col' (вертикальный) | 'pie'."""
    if kind == 'line':
        chart = LineChart()
        chart.grouping = 'standard'
        chart.smooth = True
    elif kind == 'pie':
        chart = PieChart()
        chart.dataLabels = None
    else:
        chart = BarChart()
        chart.type = kind
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height
    if x_title:
        chart.x_axis.title = x_title
    if y_title:
        chart.y_axis.title = y_title
    return chart


def color_series(chart, *colors, line=False):
    for series, color in zip(chart.series, colors):
        if line:
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 25000
            series.smooth = True
        else:
            series.graphicalProperties.solidFill = color


def _chart_from_sheet(wb, sheet, kind, title, color=None, top_n=None, **axis):
    """График по листу отчёта: категории — колонка A, значения — колонка выручки."""
    if sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    last = ws.max_row if top_n is None else min(top_n + 1, ws.max_row)
    if last <= 2:
        return None
    chart = make_chart(kind, title, width=18 if kind == 'pie' else 22, **axis)
    chart.add_data(Reference(ws, min_col=_find_rev_col(ws), min_row=1, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    if color:
        color_series(chart, color, line=(kind == 'line'))
    return chart


def add_report_charts(wb):
    """Лист «📊 Графики» основного отчёта: две колонки графиков по 25 строк."""
    ws = wb.create_sheet("📊 Графики", 0)
    ws.sheet_view.showGridLines = False
    ws['A1'] = "АНАЛИТИКА — ГРАФИКИ"
    ws['A1'].font = TITLE_FONT
    ws.row_dimensions[1].height = 30

    pairs = [
        (_chart_from_sheet(wb, '01_Месячная_статистика', 'line', "Выручка по месяцам, тыс. руб.",
                           ACCENT, x_title="Период", y_title="тыс. руб."),
         _chart_from_sheet(wb, '03_Топ_менеджеров', 'bar', "Топ-10 менеджеров, тыс. руб.",
                           ACCENT, top_n=10, x_title="тыс. руб.", y_title="Менеджер")),
        (_chart_from_sheet(wb, '02_Топ_клиентов', 'col', "Топ-10 клиентов, тыс. руб.",
                           ACCENT_2, top_n=10, y_title="тыс. руб."),
         _chart_from_sheet(wb, '04_Отрасли', 'pie', "Выручка по отраслям", top_n=10)),
    ]
    row = 3
    for left, right in pairs:
        if left is not None:
            ws.add_chart(left, f"A{row}")
        if right is not None:
            ws.add_chart(right, f"L{row + 3}")
        row += 25
